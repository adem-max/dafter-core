"""The tables of the yearly archive and the accountant's closing file
(SCF advice Q20): balance générale, grand livre, journal centralisateur,
aged balances, stock statement, fixed-asset register, TVA per month, the SCF
statements; and the writers that put them in Excel or PDF.

A table is (title, columns, rows): columns are (label, kind) with kind one of
"text", "date", "amount", "number"; rows are lists in the same order.
"""

import io
import re
from collections import OrderedDict

import frappe
from frappe import _
from frappe.utils import add_months, flt, get_first_day, getdate

from dz_core.books.export import entry_rows
from dz_core.books.journals import JOURNALS

TEXT, DATE, AMOUNT, NUMBER = "text", "date", "amount", "number"


def table(title, columns, rows):
	return frappe._dict(title=title, columns=columns, rows=rows)


def fiscal_year_start(company, date):
	from erpnext.accounts.utils import get_fiscal_year

	return getdate(get_fiscal_year(date, company=company)[1])


# ------------------------------------------------------------ balance générale


def trial_balance(company, start, end):
	"""Opening, movements and closing per account. Income and expense accounts
	open at zero on the fiscal year's first day."""
	start, end = getdate(start), getdate(end)
	year_start = fiscal_year_start(company, start)
	rows = frappe.db.sql(
		"""SELECT acc.account_number, acc.account_name, acc.root_type,
			SUM(IF(gle.posting_date < %(start)s
				AND (acc.root_type IN ('Asset', 'Liability', 'Equity') OR gle.posting_date >= %(year_start)s),
				gle.debit - gle.credit, 0)) AS opening,
			SUM(IF(gle.posting_date >= %(start)s, gle.debit, 0)) AS debit,
			SUM(IF(gle.posting_date >= %(start)s, gle.credit, 0)) AS credit
		FROM `tabGL Entry` gle JOIN `tabAccount` acc ON acc.name = gle.account
		WHERE gle.company = %(company)s AND gle.is_cancelled = 0 AND gle.posting_date <= %(end)s
		GROUP BY gle.account
		ORDER BY acc.account_number, acc.account_name""",
		{"company": company, "start": start, "end": end, "year_start": year_start},
		as_dict=True,
	)
	lines = []
	for row in rows:
		opening = flt(row.opening, 2)
		closing = flt(opening + flt(row.debit) - flt(row.credit), 2)
		if not (opening or row.debit or row.credit):
			continue
		lines.append(
			[
				row.account_number or "",
				row.account_name,
				max(opening, 0),
				max(-opening, 0),
				flt(row.debit, 2),
				flt(row.credit, 2),
				max(closing, 0),
				max(-closing, 0),
			]
		)
	if lines:
		lines.append(["", _("Total")] + [flt(sum(line[i] for line in lines), 2) for i in range(2, 8)])
	return table(
		_("Balance générale"),
		[
			(_("Compte"), TEXT),
			(_("Libellé"), TEXT),
			(_("Solde d'ouverture débit"), AMOUNT),
			(_("Solde d'ouverture crédit"), AMOUNT),
			(_("Mouvements débit"), AMOUNT),
			(_("Mouvements crédit"), AMOUNT),
			(_("Solde de clôture débit"), AMOUNT),
			(_("Solde de clôture crédit"), AMOUNT),
		],
		lines,
	)


# ------------------------------------------------------ grand livre, journaux


def general_ledger(company, start, end, entries=None):
	"""Every line by account then date, with the running balance."""
	entries = entry_rows(company, start, end) if entries is None else entries
	by_account = OrderedDict()
	for row in sorted(entries, key=lambda r: (str(r["CompteNum"]), r["EcritureDate"], r["EcritureNum"])):
		by_account.setdefault(row["CompteNum"], []).append(row)
	lines = []
	for number, rows in by_account.items():
		running = 0
		for row in rows:
			running = flt(running + row["Debit"] - row["Credit"], 2)
			lines.append(
				[
					number,
					row["CompteLib"],
					row["EcritureDate"],
					row["JournalCode"],
					row["EcritureNum"],
					row["CompAuxLib"],
					row["EcritureLib"],
					row["Debit"],
					row["Credit"],
					running,
				]
			)
	return table(
		_("Grand livre"),
		[
			(_("Compte"), TEXT),
			(_("Libellé du compte"), TEXT),
			(_("Date"), DATE),
			(_("Journal"), TEXT),
			(_("Écriture"), TEXT),
			(_("Tiers"), TEXT),
			(_("Libellé"), TEXT),
			(_("Débit"), AMOUNT),
			(_("Crédit"), AMOUNT),
			(_("Solde"), AMOUNT),
		],
		lines,
	)


def centralising_journal(company, start, end, entries=None):
	"""Debit and credit per journal and month (decree 09-110 art. 12)."""
	entries = entry_rows(company, start, end) if entries is None else entries
	totals = OrderedDict()
	for row in sorted(entries, key=lambda r: (r["EcritureDate"].strftime("%Y-%m"), r["JournalCode"])):
		key = (row["EcritureDate"].strftime("%m/%Y"), row["JournalCode"])
		debit, credit = totals.get(key, (0, 0))
		totals[key] = (debit + row["Debit"], credit + row["Credit"])
	lines = [
		[month, code, _(JOURNALS[code]), flt(debit, 2), flt(credit, 2)]
		for (month, code), (debit, credit) in totals.items()
	]
	if lines:
		lines.append(
			[
				"",
				"",
				_("Total"),
				flt(sum(line[3] for line in lines), 2),
				flt(sum(line[4] for line in lines), 2),
			]
		)
	return table(
		_("Journal centralisateur"),
		[
			(_("Mois"), TEXT),
			(_("Code"), TEXT),
			(_("Journal"), TEXT),
			(_("Débit"), AMOUNT),
			(_("Crédit"), AMOUNT),
		],
		lines,
	)


# ------------------------------------------------------------- third parties


def aged_balances(company, end, party_type):
	"""ERPNext's own "Accounts Receivable / Payable Summary" at the date."""
	if party_type == "Customer":
		from erpnext.accounts.report.accounts_receivable_summary.accounts_receivable_summary import execute

		title = _("Balance âgée clients")
	else:
		from erpnext.accounts.report.accounts_payable_summary.accounts_payable_summary import execute

		title = _("Balance âgée fournisseurs")
	filters = frappe._dict(
		company=company,
		report_date=end,
		ageing_based_on="Due Date",
		calculate_ageing_with="Report Date",
		range="30, 60, 90, 120",
	)
	columns, data = execute(filters)[:2]
	keep = [
		c for c in columns if c.get("fieldtype") in ("Currency", "Float") or c.get("fieldname") == "party"
	]
	kinds = {"Currency": AMOUNT, "Float": AMOUNT}
	return table(
		title,
		[(c.get("label"), kinds.get(c.get("fieldtype"), TEXT)) for c in keep],
		[[row.get(c["fieldname"]) for c in keep] for row in data if isinstance(row, dict)],
	)


# ---------------------------------------------------------------------- stock


def stock_statement(company, end):
	"""Quantity, moving-average cost and value per item and warehouse at the date."""
	rows = frappe.db.sql(
		"""SELECT sle.item_code, item.item_name, sle.warehouse, sle.qty_after_transaction,
			sle.valuation_rate, sle.stock_value
		FROM `tabStock Ledger Entry` sle
		JOIN (
			SELECT item_code, warehouse, MAX(posting_datetime) AS latest
			FROM `tabStock Ledger Entry`
			WHERE company = %(company)s AND is_cancelled = 0 AND posting_date <= %(end)s
			GROUP BY item_code, warehouse
		) last ON last.item_code = sle.item_code AND last.warehouse = sle.warehouse
			AND last.latest = sle.posting_datetime
		JOIN `tabItem` item ON item.name = sle.item_code
		WHERE sle.company = %(company)s AND sle.is_cancelled = 0
		ORDER BY sle.item_code, sle.warehouse, sle.creation DESC""",
		{"company": company, "end": end},
		as_dict=True,
	)
	seen, lines = set(), []
	for row in rows:  # several lines at the same time: the last one created
		key = (row.item_code, row.warehouse)
		if key in seen:
			continue
		seen.add(key)
		value = flt(
			frappe.db.sql(
				"""SELECT SUM(stock_value_difference) FROM `tabStock Ledger Entry`
				WHERE company = %s AND item_code = %s AND warehouse = %s AND is_cancelled = 0 AND posting_date <= %s""",
				(company, row.item_code, row.warehouse, end),
			)[0][0],
			2,
		)
		qty = flt(row.qty_after_transaction, 3)
		if qty or value:
			lines.append(
				[row.item_code, row.item_name, row.warehouse, qty, flt(value / qty, 2) if qty else 0, value]
			)
	if lines:
		lines.append(["", _("Total"), "", None, None, flt(sum(line[5] for line in lines), 2)])
	return table(
		_("État des stocks"),
		[
			(_("Article"), TEXT),
			(_("Désignation"), TEXT),
			(_("Entrepôt"), TEXT),
			(_("Quantité"), NUMBER),
			(_("Coût moyen (CMP)"), AMOUNT),
			(_("Valeur"), AMOUNT),
		],
		lines,
	)


def stock_gap(company, end):
	"""Stock accounts minus the stock ledger's value at the date (0 when they agree)."""
	ledger = flt(
		frappe.db.sql(
			"""SELECT SUM(stock_value_difference) FROM `tabStock Ledger Entry`
			WHERE company = %s AND is_cancelled = 0 AND posting_date <= %s""",
			(company, end),
		)[0][0]
	)
	accounts = flt(
		frappe.db.sql(
			"""SELECT SUM(gle.debit - gle.credit) FROM `tabGL Entry` gle
			JOIN `tabAccount` acc ON acc.name = gle.account
			WHERE gle.company = %s AND gle.is_cancelled = 0 AND gle.posting_date <= %s
				AND acc.account_type = 'Stock'""",
			(company, end),
		)[0][0]
	)
	return flt(accounts - ledger, 2)


# --------------------------------------------------------------- fixed assets


def asset_register(company, end):
	"""Each asset: cost, depreciation up to the date (opening and booked), net value."""
	lines = []
	for asset in frappe.get_all(
		"Asset",
		filters={"company": company, "docstatus": 1, "purchase_date": ["<=", end]},
		fields=[
			"name",
			"asset_name",
			"asset_category",
			"purchase_date",
			"available_for_use_date",
			"net_purchase_amount",
			"opening_accumulated_depreciation",
			"status",
		],
		order_by="purchase_date, name",
	):
		booked = flt(
			frappe.db.sql(
				"""SELECT SUM(ds.depreciation_amount) FROM `tabDepreciation Schedule` ds
				JOIN `tabAsset Depreciation Schedule` ads ON ads.name = ds.parent
				WHERE ads.asset = %s AND ads.docstatus = 1 AND IFNULL(ds.journal_entry, '') != ''
					AND ds.schedule_date <= %s""",
				(asset.name, end),
			)[0][0]
		)
		cost = flt(asset.net_purchase_amount, 2)
		depreciation = flt(flt(asset.opening_accumulated_depreciation) + booked, 2)
		lines.append(
			[
				asset.name,
				asset.asset_name,
				asset.asset_category,
				asset.available_for_use_date or asset.purchase_date,
				cost,
				depreciation,
				flt(cost - depreciation, 2),
				_(asset.status),
			]
		)
	return table(
		_("Registre des immobilisations"),
		[
			(_("Immobilisation"), TEXT),
			(_("Désignation"), TEXT),
			(_("Catégorie"), TEXT),
			(_("Mise en service"), DATE),
			(_("Valeur d'origine"), AMOUNT),
			(_("Amortissements cumulés"), AMOUNT),
			(_("Valeur nette"), AMOUNT),
			(_("État"), TEXT),
		],
		lines,
	)


# ------------------------------------------------------------------------ TVA


def tva_by_month(company, start, end):
	from dz_core.closing import tva_figures

	lines = []
	month = get_first_day(start)
	while month <= getdate(end):
		f = tva_figures(company, month)
		lines.append(
			[
				month.strftime("%m/%Y"),
				f.bases["19"],
				f.bases["9"],
				f.bases["0"],
				f.collected_19,
				f.collected_9,
				f.deductible,
				f.deductible_assets,
				f.to_pay,
				f.new_credit,
				f.settlement or "",
			]
		)
		month = add_months(month, 1)
	return table(
		_("TVA par mois"),
		[
			(_("Mois"), TEXT),
			(_("Base 19 %"), AMOUNT),
			(_("Base 9 %"), AMOUNT),
			(_("Exonéré"), AMOUNT),
			(_("TVA collectée 19 %"), AMOUNT),
			(_("TVA collectée 9 %"), AMOUNT),
			(_("TVA déductible"), AMOUNT),
			(_("TVA déductible immobilisations"), AMOUNT),
			(_("TVA à payer"), AMOUNT),
			(_("Crédit de TVA"), AMOUNT),
			(_("Écriture de liquidation"), TEXT),
		],
		lines,
	)


# ----------------------------------------------------------------- statements


def statement_table(company, template, start, end):
	from dz_core.books.statements import statement

	_columns, rows = statement(company, template, start, end)
	lines = []
	for row in rows:
		if row.get("is_blank_line"):
			lines.append([row.get("account") or "", None])
			continue
		value = row.get(row["_segment_info"]["period_keys"][-1]) if row.get("_segment_info") else None
		lines.append([row.get("account"), flt(value, 2)])
	return table(_(template), [(_("Rubrique"), TEXT), (_("Montant"), AMOUNT)], lines)


# -------------------------------------------------------------------- writers


def to_xlsx(tables, banner=None):
	"""One tab per table; `banner` (e.g. "BROUILLON") on top of every tab."""
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill

	book = Workbook()
	book.remove(book.active)
	for t in tables:
		# Excel refuses [ ] : * ? / \ in a sheet name, and more than 31 characters
		sheet = book.create_sheet(re.sub(r"[\[\]:*?/\\]", "-", t.title)[:31])
		if banner:
			sheet.append([banner])
			sheet["A1"].font = Font(bold=True, color="C00000")
		sheet.append([label for label, _kind in t.columns])
		header = sheet.max_row
		for cell in sheet[header]:
			cell.font = Font(bold=True)
			cell.fill = PatternFill("solid", fgColor="DDEBF7")
		for row in t.rows:
			sheet.append(list(row))
		for index, (_label, kind) in enumerate(t.columns, 1):
			fmt = {DATE: "DD/MM/YYYY", AMOUNT: "#,##0.00", NUMBER: "#,##0.###"}.get(kind)
			width = 14 if kind != TEXT else 30
			sheet.column_dimensions[sheet.cell(row=header, column=index).column_letter].width = width
			if fmt:
				for (cell,) in sheet.iter_rows(min_row=header + 1, min_col=index, max_col=index):
					cell.number_format = fmt
		sheet.freeze_panes = sheet.cell(row=header + 1, column=1)
	buffer = io.BytesIO()
	book.save(buffer)
	return buffer.getvalue()


def to_html(tables, heading, banner=None):
	def cell(value, kind):
		if value is None:
			return ""
		if kind == AMOUNT:
			return frappe.format(value, {"fieldtype": "Currency"})
		if kind == DATE:
			return frappe.format(value, {"fieldtype": "Date"})
		return frappe.utils.escape_html(str(value))

	parts = [
		"<style>body{font-family:sans-serif;font-size:9pt} table{border-collapse:collapse;width:100%;margin-bottom:18px}"
		" th,td{border:1px solid #999;padding:2px 4px} td.n{text-align:right} h2{font-size:11pt}</style>",
		f"<h1 style='font-size:13pt'>{frappe.utils.escape_html(heading)}</h1>",
	]
	if banner:
		parts.append(f"<p style='color:#c00000;font-weight:bold'>{frappe.utils.escape_html(banner)}</p>")
	for t in tables:
		parts.append(f"<h2>{frappe.utils.escape_html(t.title)}</h2><table><tr>")
		parts += [f"<th>{frappe.utils.escape_html(label)}</th>" for label, _kind in t.columns]
		parts.append("</tr>")
		for row in t.rows:
			parts.append("<tr>")
			parts += [
				f"<td class='{'n' if kind in (AMOUNT, NUMBER) else ''}'>{cell(value, kind)}</td>"
				for value, (_label, kind) in zip(row, t.columns, strict=False)
			]
			parts.append("</tr>")
		parts.append("</table>")
	return "".join(parts)


def to_pdf(tables, heading, banner=None):
	from frappe.utils.pdf import get_pdf

	return get_pdf(to_html(tables, heading, banner), {"orientation": "Landscape"})


def to_csv(t):
	"""Semicolons, decimal comma, French dates, UTF-8 with BOM: what Excel opens in Algeria."""
	import csv
	import datetime

	def text(value):
		if isinstance(value, datetime.datetime):
			return value.strftime("%d/%m/%Y %H:%M:%S")
		if isinstance(value, datetime.date):
			return value.strftime("%d/%m/%Y")
		if isinstance(value, float):
			return f"{value:.2f}".replace(".", ",")
		return value

	buffer = io.StringIO()
	writer = csv.writer(buffer, delimiter=";")
	writer.writerow([label for label, _kind in t.columns])
	for row in t.rows:
		writer.writerow([text(value) for value in row])
	return buffer.getvalue().encode("utf-8-sig")
