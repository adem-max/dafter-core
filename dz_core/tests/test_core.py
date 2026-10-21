"""dz_core on a real site: setup, tax regime, legal fields, print format, import workbook.

The first test class prepares the site once (dz_core.setup.new_client.prepare)
if it was never prepared, so the tests work on an empty test site.
"""

import os
import re
import tempfile
from html import unescape

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import random_string, today

from dz_core.dz_core.doctype.dz_core_settings.dz_core_settings import IFU, REEL
from dz_core.imports.workbook import import_workbook, make_template
from dz_core.setup.new_client import prepare
from dz_core.setup.sectors import CLOTHING, FOOD, SECTORS, create_item_attributes, create_uoms
from dz_core.setup.tax import (
	PURCHASE_TEMPLATE,
	SALES_TEMPLATE,
	STANDARD,
	apply_tax_regime,
	template_with_rate,
)

COMPANY = "Producteur Test DZ"
ABBR = "PTD"


def ensure_prepared():
	"""Prepare the site once. On a site whose setup wizard already ran (a dz_cod
	test site made by its demo loader), use the company that is there."""
	if not frappe.db.get_single_value("DZ Core Settings", "company"):
		if frappe.is_setup_complete():
			company = (
				frappe.defaults.get_global_default("company") or frappe.get_all("Company", pluck="name")[0]
			)
			prepare(company, frappe.get_cached_value("Company", company, "abbr"), FOOD, IFU)
		else:
			prepare(COMPANY, ABBR, FOOD, IFU)
	return frappe.get_doc("DZ Core Settings")


class TestPrepareGuard(IntegrationTestCase):
	def test_unknown_company_on_configured_site(self):
		ensure_prepared()
		with self.assertRaises(frappe.ValidationError):
			prepare("Société qui n'existe pas", "SQN", FOOD, IFU)


class TestSetup(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		cls.settings = ensure_prepared()
		cls.company = cls.settings.company

	def test_opening_account(self):
		from dz_core.setup.accounts import opening_account

		name = opening_account(self.company)
		self.assertEqual(frappe.db.get_value("Account", name, "account_type"), "Temporary")
		self.assertEqual(opening_account(self.company), name)

	def test_item_group_warehouses(self):
		for row in frappe.get_doc("Item Group", "Produits finis").item_group_defaults:
			if row.company == self.company:
				self.assertEqual(row.default_warehouse, self.settings.finished_goods_warehouse)
				break
		else:
			self.fail("no default warehouse on Produits finis")

	def test_sector_template(self):
		for group in SECTORS[self.settings.sector]["item_groups"]:
			self.assertTrue(frappe.db.exists("Item Group", group), group)
		for uom, _whole in SECTORS[self.settings.sector]["uoms"]:
			self.assertTrue(frappe.db.exists("UOM", uom), uom)
		for field in ("raw_materials_warehouse", "wip_warehouse", "finished_goods_warehouse"):
			warehouse = self.settings.get(field)
			self.assertEqual(frappe.db.get_value("Warehouse", warehouse, "company"), self.company)
		self.assertEqual(
			frappe.db.get_value("Company", self.company, "default_wip_warehouse"), self.settings.wip_warehouse
		)

	def test_prepare_twice_changes_nothing(self):
		# once to apply whatever a newer prepare() adds, then the real check
		prepare(self.company, ABBR, self.settings.sector, self.settings.tax_regime)
		counts = {doctype: frappe.db.count(doctype) for doctype in ("Warehouse", "Item Group", "UOM")}
		prepare(self.company, ABBR, self.settings.sector, self.settings.tax_regime)
		for doctype, count in counts.items():
			self.assertEqual(frappe.db.count(doctype), count, doctype)

	def test_tax_regimes(self):
		apply_tax_regime(self.company, IFU)
		self.assertFalse(
			frappe.get_all("Sales Taxes and Charges Template", {"company": self.company, "is_default": 1})
		)
		apply_tax_regime(self.company, REEL)
		abbr = frappe.get_cached_value("Company", self.company, "abbr")
		for doctype, title in (
			("Sales Taxes and Charges Template", SALES_TEMPLATE),
			("Purchase Taxes and Charges Template", PURCHASE_TEMPLATE),
		):
			default = frappe.get_all(doctype, {"company": self.company, "is_default": 1}, pluck="name")
			self.assertEqual(default, [f"{title} - {abbr}"])
		# every group of the company carries the standard rate
		taxes = frappe.get_doc("Item Group", "Produits finis").taxes
		self.assertIn(f"{STANDARD} - {abbr}", [row.item_tax_template for row in taxes])
		apply_tax_regime(self.company, self.settings.tax_regime)

	def test_default_print_formats(self):
		from dz_core.setup.new_client import set_default_print_formats

		set_default_print_formats()
		self.assertEqual(frappe.get_meta("Sales Invoice").default_print_format, "Facture DZ")
		self.assertEqual(frappe.get_meta("Delivery Note").default_print_format, "Bon de livraison DZ")

	def test_cancelled_invoices_can_be_printed(self):
		from dz_core.setup.new_client import allow_printing_cancelled

		allow_printing_cancelled()
		self.assertTrue(frappe.db.get_single_value("Print Settings", "allow_print_for_cancelled"))

	def test_algerian_address_prints_without_the_country(self):
		from frappe.contacts.doctype.address.address import get_address_display

		from dz_core.setup.new_client import set_address_template

		set_address_template()
		text = get_address_display(
			{
				"address_line1": "12 rue Larbi Ben M'hidi",
				"city": "Boufarik",
				"state": "Blida",
				"country": "Algeria",
			}
		)
		self.assertIn("Boufarik", text)
		self.assertIn("Blida", text)
		self.assertNotIn("Algeria", text)

	def test_wizard_item_groups_removed_unless_used(self):
		from dz_core.setup.sectors import WIZARD_ITEM_GROUPS, remove_wizard_item_groups

		remove_wizard_item_groups()
		for name in WIZARD_ITEM_GROUPS:
			if frappe.db.exists("Item Group", name):
				# kept only because an item (or a setting) still uses it
				self.assertTrue(frappe.db.exists("Item", {"item_group": name}) or self.linked(name), name)
		self.assertTrue(frappe.db.exists("Item Group", "Services"))
		for group in SECTORS[self.settings.sector]["item_groups"]:
			self.assertTrue(frappe.db.exists("Item Group", group), group)

	def linked(self, item_group):
		from frappe.desk.form.linked_with import get_linked_docs, get_linked_doctypes

		return bool(get_linked_docs("Item Group", item_group, get_linked_doctypes("Item Group")))

	def test_legal_fields(self):
		for doctype, fields in {
			"Company": ("dz_legal_form", "dz_activity", "dz_rc", "dz_nif", "dz_nis", "dz_ai", "dz_capital"),
			"Customer": ("dz_rc", "dz_nif", "dz_nis", "dz_ai"),
			"Sales Invoice": ("dz_payment_mode",),
		}.items():
			meta = frappe.get_meta(doctype)
			for field in fields:
				self.assertTrue(meta.has_field(field), f"{doctype}.{field}")


class TestWorkbook(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		cls.settings = ensure_prepared()
		# the workbook test uses the clothing sheets on whatever sector the site has
		create_item_attributes(SECTORS[CLOTHING]["item_attributes"])
		create_uoms(SECTORS[CLOTHING]["uoms"])
		cls.tag = random_string(5).upper()

	def workbook(self, fill):
		"""Make the template, let fill(wb) add rows, return the saved path."""
		from openpyxl import load_workbook

		path = os.path.join(tempfile.mkdtemp(), "import.xlsx")
		make_template(path, sector=CLOTHING)  # clothing template: has every sheet
		wb = load_workbook(path)
		fill(wb)
		wb.save(path)
		return path

	def test_errors_create_nothing(self):
		code = f"RM-{self.tag}-BAD"

		def fill(wb):
			wb["Articles"].append([code, "Sucre", "Groupe inexistant", "Kg", 120, None, "Non", "Non", None])
			wb["Stock initial"].append([code, "Entrepôt inexistant", 10, 120, None, None])
			wb["Clients"].append([f"Client {code}", "", "", "", "", "Wilaya inexistante", "", ""])

		result = import_workbook(self.workbook(fill), dry_run=0)
		self.assertEqual(len(result["errors"]), 3, result["errors"])
		self.assertFalse(frappe.db.exists("Item", code))

	def test_full_import_then_reimport(self):
		t = self.tag
		jam, sugar, fruit, jar = f"FG-{t}-CONF", f"RM-{t}-SUCRE", f"RM-{t}-FRUIT", f"PK-{t}-POT"
		model = f"FG-{t}-TSHIRT"

		def fill(wb):
			wb["Fournisseurs"].append([f"Fournisseur {t}", "0550000000", "Rue 1", "Médéa", ""])
			wb["Clients"].append(
				[
					f"Client {t}",
					"0661000000",
					"",
					"Rue 2",
					"Blida",
					"Blida",
					"000016001234567",
					"",
					"000016001234568",
					"16000000001",
				]
			)
			wb["Articles"].append(
				[sugar, "Sucre", "Matières premières", "Kg", 120, None, "Non", "Non", None, 30, 100]
			)
			wb["Articles"].append(
				[fruit, "Fraises", "Matières premières", "Kg", 300, None, "Non", "Non", None]
			)
			wb["Articles"].append([jar, "Pot 250 g", "Emballages", "Pot", 25, None, "Non", "Non", None])
			wb["Articles"].append(
				[jam, "Confiture de fraise", "Produits finis", "Pot", None, 450, "Oui", "Oui", 365]
			)
			wb["Modèles"].append([model, "T-shirt", "Produits finis", "Pièce", "M, L", "Noir, Orange", 1500])
			wb["Nomenclatures"].append([jam, 10, fruit, 1.5, "Kg"])
			wb["Nomenclatures"].append([jam, 10, sugar, 1.2, "Kg"])
			wb["Nomenclatures"].append([jam, 10, jar, 10, "Pot"])
			wb["Stock initial"].append([sugar, "Matières premières", 50, 120, None, None])
			wb["Stock initial"].append([jam, "Produits finis", 20, 300, f"L-{t}", "31/12/2027"])

		path = self.workbook(fill)
		check = import_workbook(path)
		self.assertEqual(check["errors"], [])
		self.assertFalse(frappe.db.exists("Item", jam), "a dry run must create nothing")

		result = import_workbook(path, dry_run=0)
		self.assertEqual(result["errors"], [])
		self.assertTrue(frappe.db.get_value("Item", jam, "has_batch_no"))
		reorder = frappe.db.get_value(
			"Item Reorder",
			{"parent": sugar},
			["warehouse", "warehouse_reorder_level", "warehouse_reorder_qty"],
		)
		self.assertEqual(reorder, (self.settings.raw_materials_warehouse, 30, 100))
		customer = frappe.db.get_value(
			"Customer", {"customer_name": f"Client {t}"}, ["dz_phone", "dz_wilaya", "dz_nis", "dz_ai"]
		)
		self.assertEqual(customer, ("0661000000", "09 - Blida", "000016001234568", "16000000001"))
		self.assertEqual(frappe.db.count("Item", {"variant_of": model}), 4)
		self.assertTrue(frappe.db.exists("Item", {"variant_of": model, "item_name": "T-shirt L Orange"}))
		self.assertTrue(
			frappe.db.exists("Item Attribute Value", {"parent": "Couleur", "attribute_value": "Orange"})
		)
		bom = frappe.get_doc("BOM", {"item": jam, "is_default": 1})
		self.assertEqual(len(bom.items), 3)
		self.assertEqual(bom.quantity, 10)
		self.assertTrue(frappe.db.exists("Batch", f"L-{t}"))
		self.assertEqual(
			frappe.db.get_value(
				"Bin", {"item_code": sugar, "warehouse": self.settings.raw_materials_warehouse}, "actual_qty"
			),
			50,
		)

		again = import_workbook(path, dry_run=0)
		self.assertEqual(again["created"], {})

	def test_check_finds_what_the_import_would_refuse(self):
		t = self.tag
		flour, salt, jam, jam2 = f"RM-{t}-FARINE", f"RM-{t}-SEL", f"FG-{t}-LOTA", f"FG-{t}-LOTB"
		# a lot number already used by another item on the site
		other = f"FG-{t}-AUTRE"
		frappe.get_doc(
			{
				"doctype": "Item",
				"item_code": other,
				"item_name": "Autre produit",
				"item_group": "Produits finis",
				"stock_uom": "Nos",
				"has_batch_no": 1,
			}
		).insert()
		frappe.get_doc({"doctype": "Batch", "batch_id": f"L-{t}-PRIS", "item": other}).insert()

		def fill(wb):
			wb["Clients"].append([f"Client {t} double", "", "", "", "", "", "", "", "", ""])
			wb["Clients"].append([f"client {t} DOUBLE", "", "", "", "", "", "", "", "", ""])
			# no purchase price; a price that is not a number
			wb["Articles"].append(
				[flour, "Farine", "Matières premières", "Kg", None, None, "Non", "Non", None]
			)
			wb["Articles"].append([salt, "Sel", "Matières premières", "Kg", "abc", None, "Non", "Non", None])
			wb["Articles"].append([jam, "Confiture A", "Produits finis", "Pot", 100, 450, "Oui", "Non", None])
			wb["Articles"].append(
				[jam2, "Confiture B", "Produits finis", "Pot", 100, 450, "Oui", "Non", None]
			)
			# flour is counted in Kg: nothing turns a Pot into Kg
			wb["Nomenclatures"].append([jam, 10, flour, 2, "Pot"])
			wb["Stock initial"].append([flour, "Matières premières", 10, None, None, None])  # no cost at all
			wb["Stock initial"].append([salt, "Matières premières", 5, 50, None, None])
			wb["Stock initial"].append([salt, "Matières premières", 5, 50, None, None])  # counted twice
			wb["Stock initial"].append([jam, "Produits finis", 5, 300, f"L-{t}-COMMUN", None])
			wb["Stock initial"].append([jam2, "Produits finis", 5, 300, f"L-{t}-COMMUN", None])  # same lot
			wb["Stock initial"].append([jam, "Produits finis", 5, 300, f"L-{t}-PRIS", None])  # other's lot

		result = import_workbook(self.workbook(fill), dry_run=0)
		errors = "\n".join(result["errors"])
		for expected in (
			f"client « client {t} DOUBLE » en double",
			"Prix d'achat « abc » n'est pas un nombre",
			"rien ne convertit Pot en Kg",
			f"« {flour} » : coût unitaire à remplir",
			f"« {salt} » compté deux fois",
			f"lot « L-{t}-COMMUN » : c'est déjà le lot de l'article « {jam} »",
			f"lot « L-{t}-PRIS » : c'est déjà le lot de l'article « {other} »",
		):
			self.assertIn(expected, errors)
		self.assertEqual(len(result["errors"]), 7, errors)
		self.assertFalse(frappe.db.exists("Item", jam), "nothing is created while the file has errors")

	def test_opening_stock_is_posted_100_lines_at_a_time(self):
		# Above 100 lines ERPNext submits in a background job: the importer splits
		from unittest.mock import patch

		from dz_core.imports import workbook

		codes = [f"RM-{self.tag}-N{i}" for i in range(3)]

		def fill(wb):
			for code in codes:
				wb["Articles"].append(
					[code, f"Matière {code}", "Matières premières", "Kg", 10, None, "Non", "Non"]
				)
				wb["Stock initial"].append([code, "Matières premières", 4, 10, None, None])

		with patch.object(workbook, "MAX_LINES", 2):
			result = import_workbook(self.workbook(fill), dry_run=0)
		self.assertEqual(result["errors"], [])
		self.assertEqual(result["created"]["Stock initial"], 2)
		counts = set(
			frappe.get_all("Stock Reconciliation Item", filters={"item_code": ["in", codes]}, pluck="parent")
		)
		self.assertEqual(len(counts), 2)
		for name in counts:
			self.assertEqual(frappe.db.get_value("Stock Reconciliation", name, "docstatus"), 1, name)

	def test_file_from_an_older_template_still_imports(self):
		# NIS and AI were added to "Clients" later: a file without them is fine
		name = f"Client {self.tag} ancien"

		def fill(wb):
			wb["Clients"].delete_cols(9, 2)
			wb["Clients"].append([name, "0661000001", "", "Rue 3", "Blida", "Blida", "", ""])

		result = import_workbook(self.workbook(fill), dry_run=0)
		self.assertEqual(result["errors"], [])
		self.assertTrue(frappe.db.exists("Customer", {"customer_name": name}))


def totals_of(html):
	"""The totals table of a printed Facture DZ: [(label, amount)], amounts as numbers."""
	table = re.search(r'<table class="dz-totals">(.*?)</table>', html, re.S).group(1)
	rows = re.findall(r'<tr[^>]*><td>(.*?)</td><td class="num">(.*?)</td></tr>', table, re.S)
	return [(unescape(label).strip(), number(amount)) for label, amount in rows]


def number(text):
	"""'- 2 142,00' -> -2142.0 (montant() writes French separators)."""
	text = "".join(unescape(text).split()).replace(",", ".")
	return -float(text[1:]) if text.startswith("-") else float(text)


class TestPrintFormat(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		cls.settings = ensure_prepared()
		cls.code = f"PF-{random_string(5).upper()}"
		frappe.get_doc(
			{
				"doctype": "Item",
				"item_code": cls.code,
				"item_name": "Produit test",
				"item_group": "Produits finis",
				"stock_uom": "Nos",
				"is_stock_item": 0,
			}
		).insert()
		cls.customer = (
			frappe.get_doc(
				{
					"doctype": "Customer",
					"customer_name": f"Client facture {cls.code}",
					"dz_nif": "NIF-CLIENT-1",
				}
			)
			.insert()
			.name
		)

	def invoice(self, **values):
		return frappe.get_doc(
			{
				"doctype": "Sales Invoice",
				"company": self.settings.company,
				"customer": self.customer,
				"posting_date": today(),
				"due_date": today(),
				"dz_payment_mode": "Espèces",
				"items": [{"item_code": self.code, "qty": 2, "rate": 1000}],
				**values,
			}
		).insert()

	def test_invoice_renders_legal_mentions(self):
		company = frappe.get_doc("Company", self.settings.company)
		company.update({"dz_legal_form": "Auto-entrepreneur", "dz_nif": "NIF-TEST-123", "dz_rc": "AE-TEST-9"})
		company.save()
		invoice = self.invoice()
		html = frappe.get_print("Sales Invoice", invoice.name, "Facture DZ", doc=invoice)
		for text in ("FACTURE", "NIF-TEST-123", "AE-TEST-9", "NIF-CLIENT-1", "Espèces", "BROUILLON"):
			self.assertIn(text, html)
		if self.settings.invoice_note:
			self.assertIn(self.settings.invoice_note.replace("'", "&#39;"), html.replace("'", "&#39;"))

	def test_totals_add_up_with_a_discount_and_tva(self):
		from erpnext.controllers.accounts_controller import get_taxes_and_charges

		template = template_with_rate("Sales Taxes and Charges Template", self.settings.company, 19)
		if not template:
			self.skipTest("no 19 % template on this site")
		taxes = get_taxes_and_charges("Sales Taxes and Charges Template", template)
		for apply_on in ("Net Total", "Grand Total"):
			with self.subTest(apply_on):
				invoice = self.invoice(
					taxes_and_charges=template, taxes=taxes, apply_discount_on=apply_on, discount_amount=200
				)
				rows = totals_of(frappe.get_print("Sales Invoice", invoice.name, "Facture DZ", doc=invoice))
				labels = [label for label, _amount in rows]
				amounts = dict(rows)
				self.assertEqual(labels[:3], ["Total HT (DA)", "Remise", "Total HT après remise (DA)"])
				self.assertEqual(amounts["Total HT (DA)"], 2000)  # what the lines add up to
				# each line of the table follows from the one above
				self.assertAlmostEqual(
					amounts["Total HT (DA)"] + amounts["Remise"], amounts["Total HT après remise (DA)"], 2
				)
				self.assertAlmostEqual(amounts["Total HT après remise (DA)"], invoice.net_total, 2)
				ttc = amounts["Total TTC (DA)"]
				self.assertAlmostEqual(sum(amount for _label, amount in rows[2:-1]), ttc, 2)
				self.assertAlmostEqual(ttc, invoice.rounded_total or invoice.grand_total, 2)

	def test_credit_note_prints_as_an_avoir(self):
		from erpnext.accounts.doctype.sales_invoice.sales_invoice import make_sales_return

		invoice = self.invoice()
		invoice.submit()
		credit_note = make_sales_return(invoice.name)
		credit_note.insert()
		html = frappe.get_print("Sales Invoice", credit_note.name, "Facture DZ", doc=credit_note)
		text = unescape(html)
		self.assertIn("FACTURE D'AVOIR", text)
		self.assertIn(f"Facture d'origine : {invoice.name}", text)
		self.assertIn("Arrêté le présent avoir à la somme de : deux mille dinars algériens.", text)
		self.assertNotIn("moins", text)
		# amounts printed positive: the title says it is a credit
		self.assertEqual(totals_of(html), [("Total HT (DA)", 2000), ("Total de l'avoir (DA)", 2000)])


class TestAmounts(IntegrationTestCase):
	def test_amount_in_french_words(self):
		from dz_core.utils import montant, montant_en_lettres

		self.assertEqual(montant_en_lettres(32400), "trente-deux mille quatre cents dinars algériens")
		self.assertEqual(montant_en_lettres(1), "un dinar algérien")
		self.assertEqual(
			montant_en_lettres(2000000.5), "deux millions de dinars algériens et cinquante centimes"
		)
		# a credit note's total: "moins" once, plural as for a positive amount
		self.assertEqual(
			montant_en_lettres(-32400.5),
			"moins trente-deux mille quatre cents dinars algériens et cinquante centimes",
		)
		self.assertEqual(montant(32400), "32 400,00")
