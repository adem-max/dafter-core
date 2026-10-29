"""Opening of the books from the accountant's trial balance (dz_core/imports/opening.py).

On 1 January 2020, a day nothing else posts on, so the opening account starts empty.

Run: bench --site <site> run-tests --module dz_core.tests.test_opening
"""

import os
import tempfile

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import flt, getdate

from dz_core.imports import opening
from dz_core.imports.workbook import Report
from dz_core.setup.accounts import account_by_number, opening_account
from dz_core.setup.scf import apply_scf_setup
from dz_core.tests.test_core import ensure_prepared

DAY = getdate("2020-01-01")


def party(doctype, name, **values):
	field = frappe.scrub(doctype) + "_name"
	return frappe.db.get_value(doctype, {field: name}) or (
		frappe.get_doc({"doctype": doctype, field: name, **values}).insert().name
	)


class TestOpening(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		cls.company = ensure_prepared().company
		apply_scf_setup(cls.company)
		if not frappe.db.exists(
			"Fiscal Year", {"year_start_date": ["<=", DAY], "year_end_date": [">=", DAY]}
		):
			frappe.get_doc(
				{
					"doctype": "Fiscal Year",
					"year": "2020",
					"year_start_date": DAY,
					"year_end_date": "2020-12-31",
				}
			).insert()
		cls.customer = party("Customer", "Client ouverture test", customer_type="Company")
		cls.supplier = party(
			"Supplier", "Fournisseur ouverture test", supplier_group="Fournisseurs de stocks"
		)

	def workbook(self, balance, customers=(), suppliers=()):
		from openpyxl import Workbook

		wb = Workbook()
		wb.remove(wb.active)
		for name, rows in (
			(opening.BALANCE, balance),
			(opening.CUSTOMERS, customers),
			(opening.SUPPLIERS, suppliers),
		):
			ws = wb.create_sheet(name)
			ws.append([header for header, _key, _required in opening.SHEETS[name]])
			for row in rows:
				ws.append(list(row))
		handle, path = tempfile.mkstemp(suffix=".xlsx")
		os.close(handle)
		wb.save(path)
		self.addCleanup(os.remove, path)
		return path

	def balanced(self):
		return self.workbook(
			[
				("5121", "Banque", 10000, 0),
				("4111", "Clients", 3000, 0),
				("40132", "Fournisseurs", 0, 2000),
				("108", "Exploitant", 0, 11000),
			],
			customers=[("Client ouverture test", "F-118", "15/12/2019", "15/01/2020", 3000)],
			suppliers=[("Fournisseur ouverture test", "FA-7", "20/12/2019", "20/01/2020", 2000)],
		)

	def checked(self, path):
		from openpyxl import load_workbook

		report = Report()
		wb = load_workbook(path, data_only=True)
		data = {name: opening.read_rows(wb, name, report) for name in opening.SHEETS}
		return opening.check(self.company, DAY, data, report), report

	def test_a_balanced_opening_posts_and_empties_479(self):
		result = opening.import_opening(self.balanced(), DAY, dry_run=1)
		self.assertEqual(result["errors"], [])
		plan, report = self.checked(self.balanced())
		opening.post(self.company, DAY, plan, report)
		self.assertEqual(report.created.keys(), {"Sales Invoice", "Purchase Invoice", "Journal Entry"})
		self.assertAlmostEqual(opening.posted(self.company, opening_account(self.company), DAY), 0, 2)
		bank = account_by_number(self.company, "5121")
		moved = flt(
			frappe.db.sql(
				"SELECT SUM(debit - credit) FROM `tabGL Entry` WHERE account = %s AND posting_date = %s AND is_cancelled = 0",
				(bank, DAY),
			)[0][0]
		)
		self.assertGreaterEqual(moved, 10000)
		invoice = report.created["Sales Invoice"][0]
		self.assertEqual(
			frappe.db.get_value("Sales Invoice", invoice, ["is_opening", "outstanding_amount"]), ("Yes", 3000)
		)

	def test_errors_are_listed(self):
		path = self.workbook(
			[
				("5121", "Banque", 10000, 0),
				("999999", "Inconnu", 0, 5),
				("4111", "Clients", 3000, 0),
				("311", "Matières premières", 500, 0),  # never counted
			],
			customers=[("Client ouverture test", "F-1", "01/12/2019", "", 2500)],
		)
		result = opening.import_opening(path, DAY, dry_run=1)
		errors = " | ".join(result["errors"])
		self.assertIn("999999", errors)
		self.assertIn("pas équilibrée", errors)
		self.assertIn("total des factures", errors)
		self.assertIn("stock 311", errors)

	def advances(self, with_parties=True):
		return self.workbook(
			[
				("5121", "Banque", 10000, 0, None),
				("4191", "Avances reçues", 0, 1500, "Client ouverture test" if with_parties else None),
				("40913", "Avances versées", 800, 0, "Fournisseur ouverture test" if with_parties else None),
				("108", "Exploitant", 0, 9300, None),
			]
		)

	def test_advances_go_to_the_opening_entry_with_their_party(self):
		"""4191 and 40913 are receivable / payable but hold no invoice (review of 27 Oct, 2.2)."""
		plan, report = self.checked(self.advances())
		errors = [e for e in report.errors if "déjà ouverts" not in e]
		self.assertEqual(errors, [])
		parties = {line.number: line.party for line in plan["journal"]}
		self.assertEqual(parties["4191"], self.customer)
		self.assertEqual(parties["40913"], self.supplier)
		report = Report()
		opening.post(self.company, DAY, plan, report)
		entry = frappe.get_doc("Journal Entry", report.created["Journal Entry"][0])
		advance = next(
			row for row in entry.accounts if row.account == account_by_number(self.company, "4191")
		)
		self.assertEqual(
			(advance.party_type, advance.party, advance.credit), ("Customer", self.customer, 1500)
		)
		self.assertAlmostEqual(opening.posted(self.company, opening_account(self.company), DAY), 0, 2)

	def test_an_advance_without_its_party_is_refused(self):
		_plan, report = self.checked(self.advances(with_parties=False))
		errors = " | ".join(report.errors)
		self.assertIn("compte de tiers 4191", errors)
		self.assertIn("compte de tiers 40913", errors)

	def test_a_file_without_the_tiers_column_still_reads(self):
		from openpyxl import Workbook, load_workbook

		wb = Workbook()
		ws = wb.active
		ws.title = opening.BALANCE
		ws.append(["Compte", "Libellé", "Débit", "Crédit"])
		ws.append(["5121", "Banque", 100, 0])
		handle, path = tempfile.mkstemp(suffix=".xlsx")
		os.close(handle)
		wb.save(path)
		self.addCleanup(os.remove, path)
		report = Report()
		rows = opening.read_rows(load_workbook(path), opening.BALANCE, report)
		self.assertEqual(report.errors, [])
		self.assertEqual(rows[0][1]["party"], None)

	def test_a_second_opening_is_refused(self):
		"""Run again, the import would post every balance twice, and 479 would still end at zero."""
		if not opening.already_opened(self.company):
			plan, report = self.checked(self.balanced())
			opening.post(self.company, DAY, plan, report)
		self.assertTrue(opening.already_opened(self.company))
		# a dry run: the real one commits, and stops on the same errors before posting
		result = opening.import_opening(self.balanced(), DAY, dry_run=1)
		self.assertIn("déjà ouverts", " | ".join(result["errors"]))
