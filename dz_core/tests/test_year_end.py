"""SCF statements, year end, closing file and yearly archive (dz_core/books/).

The year is never closed here: a Period Closing Voucher on a test site would
stop every later test from posting in that year. The closing is tested up to
its refusals and its choice of account.

Run: bench --site <site> run-tests --module dz_core.tests.test_year_end
"""

import io

import frappe
from erpnext.accounts.utils import get_fiscal_year
from frappe.tests import IntegrationTestCase
from frappe.utils import today

from dz_core.books import archive as archives
from dz_core.books import statements as st
from dz_core.books.closing_file import make_closing_file
from dz_core.books.year_end import close_year, closing_account, year_bounds, year_checks, year_result
from dz_core.setup.accounts import account_by_number
from dz_core.setup.scf import apply_scf_setup
from dz_core.tests.test_core import ensure_prepared


class TestYearEnd(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		cls.company = ensure_prepared().company
		apply_scf_setup(cls.company)
		st.setup_statements()
		cls.fiscal_year = get_fiscal_year(today(), company=cls.company)[0]
		cls.start, cls.end = year_bounds(cls.fiscal_year)

	def journal(self, debit, credit, amount):
		entry = frappe.get_doc(
			{
				"doctype": "Journal Entry",
				"company": self.company,
				"posting_date": today(),
				"accounts": [
					{"account": account_by_number(self.company, debit), "debit_in_account_currency": amount},
					{
						"account": account_by_number(self.company, credit),
						"credit_in_account_currency": amount,
					},
				],
			}
		).insert()
		entry.submit()
		return entry

	def test_statements_balance_and_give_the_ledger_result(self):
		self.journal("5121", "7011", 1000)
		self.journal("6132", "5121", 300)
		income = st.amounts(self.company, st.INCOME_STATEMENT, self.start, self.end)
		self.assertAlmostEqual(income["RNET"], year_result(self.company, self.start, self.end), 2)
		sheet = st.amounts(self.company, st.BALANCE_SHEET, self.start, self.end)
		self.assertAlmostEqual(sheet["TA"], sheet["TP"], 2)

	def test_the_result_goes_to_120_or_129(self):
		self.assertIn("120", closing_account(self.company, 10))
		self.assertIn("129", closing_account(self.company, -10))

	def test_closing_is_refused_while_a_check_fails(self):
		self.journal("6589", "5121", 75)  # left to reclassify
		rows = year_checks(self.company, self.fiscal_year)
		self.assertEqual(next(r for r in rows if r["code"] == "6589")["status"], "À voir")
		with self.assertRaises(frappe.ValidationError):
			close_year(self.company, self.fiscal_year)

	def test_the_archive_needs_a_frozen_year(self):
		with self.assertRaises(frappe.ValidationError):
			archives.make_archive(self.company, self.fiscal_year)

	def test_closing_file_refused_then_made_as_a_draft(self):
		from openpyxl import load_workbook

		self.journal("6589", "5121", 75)
		with self.assertRaises(frappe.ValidationError):
			make_closing_file(self.company, self.fiscal_year)
		name = make_closing_file(self.company, self.fiscal_year, draft=1)
		archive = frappe.get_doc("DZ Archive", name)
		self.assertEqual((archive.kind, archive.draft), (archives.CLOSING_FILE, 1))
		self.assertEqual([file for file, ok in archives.verify(name) if not ok], [])
		workbook = next(
			row for row in archive.files if row.file_name.endswith(".xlsx") and "cloture" in row.file_name
		)
		content = frappe.get_doc("File", {"file_url": workbook.file, "attached_to_name": name}).get_content()
		book = load_workbook(io.BytesIO(content))
		for sheet in (
			"Contrôles",
			"Balance générale",
			"Grand livre",
			"Journal centralisateur",
			st.BALANCE_SHEET,
		):
			self.assertIn(sheet, book.sheetnames)
		self.assertIn("BROUILLON", str(book["Contrôles"]["A1"].value))


class TestAfterTheClose(IntegrationTestCase):
	"""Once a year is closed, its closing voucher empties classes 6 and 7 on the
	year-end date. The year's own figures must not move (review of 27 October
	2026, 1.1: the class 72 check turned "À voir" and blocked the closing
	file, and the IFU booked read 0).

	No real closing voucher here: it would block posting in that year on this
	test site for good (CLAUDE.md). The voucher's ledger lines are written by
	hand, in this test's transaction, which is rolled back."""

	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		cls.company = ensure_prepared().company
		apply_scf_setup(cls.company)
		cls.fiscal_year = get_fiscal_year(today(), company=cls.company)[0]
		cls.end = year_bounds(cls.fiscal_year)[1]

	def line(self, voucher_type, voucher_no, number_or_account, debit=0, credit=0, posting_date=None):
		from dz_core.closing import accounts_by_prefix

		if frappe.db.exists("Account", number_or_account):
			account = number_or_account
		else:
			account = (
				account_by_number(self.company, number_or_account)
				or accounts_by_prefix(self.company, number_or_account)[0]
			)
		gle = frappe.new_doc("GL Entry")
		gle.update(
			{
				"name": frappe.generate_hash(length=12),
				"company": self.company,
				"account": account,
				"debit": debit,
				"credit": credit,
				"debit_in_account_currency": debit,
				"credit_in_account_currency": credit,
				"posting_date": posting_date or today(),
				"voucher_type": voucher_type,
				"voucher_no": voucher_no,
				"is_cancelled": 0,
				"fiscal_year": self.fiscal_year,
			}
		)
		gle.db_insert()
		return account

	def test_the_years_figures_do_not_move_when_it_is_closed(self):
		from dz_core.books.ifu import year_figures
		from dz_core.closing import closing_rows

		# a manufacture (finished goods in, production stocked) and the IFU, this year
		self.line("Stock Entry", "TEST-MANUF-1", "35", debit=1000)
		production = self.line("Stock Entry", "TEST-MANUF-1", "72", credit=1000)
		self.line("Journal Entry", "TEST-IFU-1", "6428", debit=500)
		self.line("Journal Entry", "TEST-IFU-1", "4477", credit=500)

		def row_72():
			return next(r for r in closing_rows(self.company, self.end) if r["code"] == "72")

		before = row_72()
		booked = year_figures(self.company, self.fiscal_year).booked

		# the closing voucher's lines: 72 and 6428 emptied into the result, on the year end
		self.line("Period Closing Voucher", "TEST-PCV-1", production, debit=1000, posting_date=self.end)
		self.line("Period Closing Voucher", "TEST-PCV-1", "6428", credit=500, posting_date=self.end)
		self.line(
			"Period Closing Voucher",
			"TEST-PCV-1",
			closing_account(self.company, 500),
			credit=500,
			posting_date=self.end,
		)

		after = row_72()
		self.assertEqual((after["amount"], after["status"]), (before["amount"], before["status"]))
		self.assertEqual(year_figures(self.company, self.fiscal_year).booked, booked)
		self.assertGreaterEqual(booked, 500)
