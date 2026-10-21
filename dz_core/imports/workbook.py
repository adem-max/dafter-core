"""The import workbook: the client fills one Excel file, we import it in one go.

1. Make the empty workbook for a client's sector and send it to them:

    bench --site client1.localhost execute dz_core.imports.workbook.make_template \\
        --kwargs "{'path': '/tmp/import-client1.xlsx'}"

2. When it comes back, check it without creating anything (dry run, the default):

    bench --site client1.localhost execute dz_core.imports.workbook.import_workbook \\
        --kwargs "{'path': '/tmp/import-client1.xlsx'}"

3. When the check shows no errors, import for real:

    bench --site client1.localhost execute dz_core.imports.workbook.import_workbook \\
        --kwargs "{'path': '/tmp/import-client1.xlsx', 'dry_run': 0}"

Rules:
- Rows whose first cell starts with "EXEMPLE" are skipped (the template has one per sheet).
- Anything that already exists (same code or name) is skipped, never changed:
  importing the same file twice creates nothing new.
- Nothing is created while any row has an error: fix the file, check again.
- Order: suppliers, customers, items, models (clothing variants), BOMs, opening stock.
- Opening stock can come in several files (raw materials one day, finished
  goods the next), but an item that already has stock in a warehouse is
  refused there: after go-live, counts are corrected in ERPNext itself.
- A lot number belongs to one item: ERPNext names every batch by its number
  alone, so two products cannot share "L2026-10".
"""

import datetime
import json

import frappe
from frappe.utils import cint, flt, getdate

from dz_core.dz_core.doctype.dz_core_settings.dz_core_settings import get_settings
from dz_core.phone import normalize_phone
from dz_core.setup.sectors import CLOTHING, FINISHED_GROUPS, SECTORS
from dz_core.setup.wilayas import find_wilaya

EXAMPLE = "EXEMPLE"
YES = {"oui", "o", "yes", "y", "1", "x"}
NO = {"non", "n", "no", "0", ""}
# ERPNext submits a Stock Reconciliation of more than 100 lines in a background
# job, and a failure there goes unnoticed: the opening stock is posted 100 lines
# at a time
MAX_LINES = 100

# Sheet name -> columns (header text, key, required)
SHEETS = {
	"Fournisseurs": [
		("Nom", "name", True),
		("Téléphone", "phone", False),
		("Adresse", "address", False),
		("Ville", "city", False),
		("NIF", "nif", False),
	],
	"Clients": [
		("Nom", "name", True),
		("Téléphone", "phone", False),
		("E-mail", "email", False),
		("Adresse", "address", False),
		("Ville", "city", False),
		("Wilaya", "wilaya", False),
		("NIF", "nif", False),
		("RC", "rc", False),
		# the buyer's NIS and tax article are legal mentions too (decree 05-468)
		("NIS", "nis", False),
		("Article d'imposition (AI)", "ai", False),
	],
	"Articles": [
		("Code", "code", True),
		("Désignation", "name", True),
		("Groupe", "group", True),
		("Unité", "uom", True),
		("Prix d'achat", "buying", False),
		("Prix de vente", "selling", False),
		("Suivi par lot (Oui/Non)", "batch", False),
		("Date d'expiration (Oui/Non)", "expiry", False),
		("Durée de conservation (jours)", "shelf_life", False),
		# below this stock, the item shows in "A acheter" and on the home page
		("Seuil de réapprovisionnement", "reorder_level", False),
		("Quantité à commander", "reorder_qty", False),
	],
	"Modèles": [
		("Code modèle", "code", True),
		("Désignation", "name", True),
		("Groupe", "group", True),
		("Unité", "uom", True),
		("Tailles (séparées par des virgules)", "sizes", True),
		("Couleurs (séparées par des virgules)", "colours", True),
		("Prix de vente", "selling", False),
	],
	"Nomenclatures": [
		("Produit fini (code)", "product", True),
		("Quantité produite", "product_qty", True),
		("Composant (code)", "component", True),
		("Quantité", "qty", True),
		("Unité", "uom", False),
	],
	"Stock initial": [
		("Code article", "code", True),
		("Entrepôt", "warehouse", True),
		("Quantité", "qty", True),
		("Coût unitaire", "rate", False),
		("Lot", "batch", False),
		("Date d'expiration (JJ/MM/AAAA)", "expiry", False),
	],
}

EXAMPLES = {
	"Fournisseurs": ["EXEMPLE Apiculteurs de Médéa", "0550 00 00 00", "Rue des Oliviers", "Médéa", ""],
	"Clients": [
		"EXEMPLE Épicerie El Baraka",
		"0661 00 00 00",
		"",
		"12 rue Larbi Ben M'hidi",
		"Blida",
		"Blida",
		"",
		"",
		"",
		"",
	],
	"Articles": [
		"EXEMPLE-MIEL-250",
		"Miel de jujubier 250 g",
		"Produits finis",
		"Pot",
		0,
		1800,
		"Oui",
		"Oui",
		730,
		"",
		"",
	],
	"Modèles": [
		"EXEMPLE-TSHIRT",
		"T-shirt coton",
		"Produits finis",
		"Pièce",
		"S, M, L, XL",
		"Noir, Blanc",
		1500,
	],
	"Nomenclatures": ["EXEMPLE-MIEL-250", 1, "EXEMPLE-POT-250", 1, "Pot"],
	"Stock initial": ["EXEMPLE-MIEL-250", "Produits finis", 40, 1100, "L2026-10", "31/10/2028"],
}

README = [
	"Comment remplir ce fichier",
	"",
	"Une feuille par type de données. Ne changez pas les titres des colonnes.",
	"Chaque feuille a une ligne d'exemple (code commençant par EXEMPLE) : laissez-la ou supprimez-la, elle n'est pas importée.",
	"Groupes possibles : {groups}.",
	"Unités : celles d'ERPNext (Kg, Gram, Litre, Millilitre, Nos, Meter...) et {uoms}.",
	"Entrepôts (Stock initial) : Matières premières, Travaux en cours, Produits finis.",
	"Nomenclatures : une ligne par composant. Exemple : pour 10 pots de confiture, 5 Kg de fruits et 4 Kg de sucre = deux lignes avec Quantité produite = 10.",
	"Stock initial : comptez réellement votre stock avant de remplir. Un chiffre faux rend tous les rapports faux.",
	"Seuil de réapprovisionnement : quand le stock d'un article passe en dessous, il apparaît dans « À acheter ». Quantité à commander : ce que vous achetez d'habitude d'un coup.",
	"Oui/Non : écrivez Oui ou Non.",
]


# ---------------------------------------------------------------- template


def make_template(path, sector=None):
	"""Write the empty workbook for this site's sector (or the one given)."""
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill

	sector = sector or get_settings().sector
	template = SECTORS[sector]
	wb = Workbook()
	readme = wb.active
	readme.title = "Lisez-moi"
	groups = ", ".join(template["item_groups"])
	uoms = ", ".join(name for name, _whole in template["uoms"])
	for line in README:
		readme.append([line.format(groups=groups, uoms=uoms)])
	readme["A1"].font = Font(bold=True, size=14)
	readme.column_dimensions["A"].width = 120

	for sheet_name, columns in SHEETS.items():
		if sheet_name == "Modèles" and sector != CLOTHING:
			continue
		ws = wb.create_sheet(sheet_name)
		ws.append([header for header, _key, _req in columns])
		for cell, (_header, _key, required) in zip(ws[1], columns, strict=True):
			cell.font = Font(bold=True)
			if required:
				cell.fill = PatternFill("solid", fgColor="FFF2CC")
		ws.append(EXAMPLES[sheet_name])
		for cell in ws[2]:
			cell.font = Font(italic=True, color="808080")
		for index, (header, _key, _req) in enumerate(columns):
			ws.column_dimensions[chr(ord("A") + index)].width = max(14, len(header) + 2)
		ws.freeze_panes = "A2"
	wb.save(path)
	print(f"Workbook written: {path} ({sector}). Yellow columns are required.")
	return path


# ---------------------------------------------------------------- reading


class Report:
	def __init__(self):
		self.errors = []
		self.created = {}
		self.skipped = {}

	def error(self, sheet, row, message):
		self.errors.append(f"{sheet}, ligne {row} : {message}")

	def done(self, kind, name):
		self.created.setdefault(kind, []).append(name)

	def skip(self, kind, name):
		self.skipped.setdefault(kind, []).append(name)

	def as_dict(self):
		return {
			"errors": self.errors,
			"created": {kind: len(names) for kind, names in self.created.items()},
			"skipped (already there)": {kind: len(names) for kind, names in self.skipped.items()},
		}


def resolve_path(path):
	"""A server path, or a File URL (/files/... or /private/files/...) uploaded in ERPNext."""
	if path.startswith(("/files/", "/private/files/")):
		return frappe.get_doc("File", {"file_url": path}).get_full_path()
	return path


def read_rows(wb, sheet_name, report):
	"""[(excel row number, {key: value})] of a sheet, examples and blank rows left out."""
	if sheet_name not in wb.sheetnames:
		return []
	columns = SHEETS[sheet_name]
	ws = wb[sheet_name]
	header = [str(cell.value or "").strip() for cell in ws[1]]
	positions = {}
	for text, key, required in columns:
		if text not in header:
			if not required:
				continue  # an optional column added after the client's template was made
			report.error(sheet_name, 1, f"colonne « {text} » introuvable (ne changez pas les titres)")
			return []
		positions[key] = header.index(text)

	rows = []
	for number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
		row = {key: None for _text, key, _required in columns}
		row.update(
			{key: clean(values[index] if index < len(values) else None) for key, index in positions.items()}
		)
		if not any(value not in (None, "") for value in row.values()):
			continue
		if str(next(iter(row.values())) or "").upper().startswith(EXAMPLE):
			continue
		missing = [text for text, key, required in columns if required and row[key] in (None, "")]
		if missing:
			report.error(sheet_name, number, "à remplir : " + ", ".join(missing))
			continue
		rows.append((number, row))
	return rows


def clean(value):
	if isinstance(value, str):
		value = " ".join(value.split())
		return value
	return value


def yes_no(value, sheet, number, report, default=False):
	if value is None:
		return default
	text = str(value).strip().lower()
	if text in YES:
		return True
	if text in NO:
		return default if text == "" else False
	report.error(sheet, number, f"« {value} » : écrivez Oui ou Non")
	return default


def as_date(value, sheet, number, report):
	if value in (None, ""):
		return None
	if isinstance(value, datetime.datetime | datetime.date):
		return getdate(value)
	try:
		return datetime.datetime.strptime(str(value).strip(), "%d/%m/%Y").date()
	except ValueError:
		report.error(sheet, number, f"date « {value} » : format attendu JJ/MM/AAAA")
		return None


def as_number(value, sheet, number, report, label, positive=True):
	"""A cell as a number: 1800, "1 800,50" or "1800.5". Text that is not a
	number is an error (flt() alone would quietly turn it into 0)."""
	if isinstance(value, str):
		try:
			result = float("".join(value.split()).replace(",", "."))  # any space, also non-breaking
		except ValueError:
			report.error(sheet, number, f"{label} « {value} » n'est pas un nombre")
			return 0
	else:
		result = flt(value)
	if positive and result <= 0:
		report.error(sheet, number, f"{label} doit être supérieur à 0")
	return result


def split_list(value):
	return [part.strip() for part in str(value or "").split(",") if part.strip()]


def unit_converts(item_code, uom, stock_uom):
	"""Can ERPNext turn `uom` into the item's stock unit: the same unit, a
	conversion on the item (or its model), or a general one (Kg and Gram)?
	Without one, a recipe line silently counts 1 for 1."""
	from erpnext.stock.doctype.item.item import get_uom_conv_factor

	if not stock_uom or uom == stock_uom:
		return True
	items = [item_code, frappe.db.get_value("Item", item_code, "variant_of")]
	if frappe.db.exists(
		"UOM Conversion Detail",
		{"parenttype": "Item", "parent": ["in", [code for code in items if code]], "uom": uom},
	):
		return True
	return bool(get_uom_conv_factor(uom, stock_uom))


# ---------------------------------------------------------------- import


def import_workbook(path, dry_run=1, opening_date=None):
	"""Check the workbook; with dry_run=0 and no errors, create everything.

	opening_date (YYYY-MM-DD, optional): the day the opening stock was counted.
	Default today; a past date lets later movements (or a demo's history) follow it."""
	from openpyxl import load_workbook

	dry_run = cint(dry_run)
	settings = get_settings()
	if not settings.company:
		frappe.throw("Remplissez d'abord DZ Core Settings (dz_core.setup.new_client.prepare).")

	report = Report()
	wb = load_workbook(resolve_path(path), data_only=True)
	data = {sheet: read_rows(wb, sheet, report) for sheet in SHEETS}

	plan = check(data, settings, report)
	if report.errors or dry_run:
		result = report.as_dict()
		result["mode"] = "vérification seulement" if dry_run else "rien n'a été créé"
		result["would create"] = {kind: len(items) for kind, items in plan.items()}
		print(json.dumps(result, indent=1, ensure_ascii=False))
		return result

	create(plan, settings, report, opening_date)
	frappe.db.commit()
	result = report.as_dict()
	print(json.dumps(result, indent=1, ensure_ascii=False))
	return result


def check(data, settings, report):
	"""Validate every row against the site and the rest of the workbook.
	Returns what would be created, per kind, without creating anything."""
	template = SECTORS[settings.sector]
	plan = {"suppliers": [], "customers": [], "items": [], "models": [], "boms": [], "stock": []}
	known_uoms = set(frappe.get_all("UOM", pluck="name"))
	known_groups = set(frappe.get_all("Item Group", pluck="name"))

	names = set()
	for number, row in data["Fournisseurs"]:
		if twice(names, row["name"]):
			report.error("Fournisseurs", number, f"fournisseur « {row['name']} » en double dans le fichier")
		elif frappe.db.exists("Supplier", {"supplier_name": row["name"]}):
			report.skip("Fournisseur", row["name"])
		else:
			plan["suppliers"].append(row)

	names = set()
	for number, row in data["Clients"]:
		if twice(names, row["name"]):
			report.error("Clients", number, f"client « {row['name']} » en double dans le fichier")
			continue
		if row["address"] and not row["city"]:
			report.error("Clients", number, "Ville à remplir quand l'adresse est remplie")
		if row["wilaya"] and not find_wilaya(row["wilaya"]):
			report.error("Clients", number, f"Wilaya « {row['wilaya']} » inconnue (ex. 16, Alger)")
		if frappe.db.exists("Customer", {"customer_name": row["name"]}):
			report.skip("Client", row["name"])
		else:
			plan["customers"].append(row)

	batches_allowed = cint(
		frappe.db.get_single_value("Stock Settings", "enable_serial_and_batch_no_for_item")
	)
	codes_in_file = set()
	units_in_file = {}  # item code -> stock unit, for the recipe lines
	for sheet, kind in (("Articles", "items"), ("Modèles", "models")):
		for number, row in data[sheet]:
			code = str(row["code"])
			if code in codes_in_file:
				report.error(sheet, number, f"code « {code} » en double dans le fichier")
				continue
			codes_in_file.add(code)
			units_in_file[code] = row["uom"]
			if row["group"] not in known_groups:
				report.error(sheet, number, f"groupe « {row['group']} » inconnu")
			if row["uom"] not in known_uoms:
				report.error(sheet, number, f"unité « {row['uom']} » inconnue")
			if kind == "items":
				in_batch_group = row["group"] in template["batch_groups"]
				row["batch"] = yes_no(row["batch"], sheet, number, report, default=in_batch_group)
				row["expiry"] = yes_no(row["expiry"], sheet, number, report, default=row["batch"])
				if row["expiry"] and not row["batch"]:
					report.error(sheet, number, "une date d'expiration demande le suivi par lot")
				if row["batch"] and not batches_allowed:
					report.error(
						sheet,
						number,
						"suivi par lot désactivé sur ce site (Stock Settings, Activate Serial / Batch No for Item)",
					)
				for key, label in (
					("buying", "Prix d'achat"),
					("selling", "Prix de vente"),
					("shelf_life", "Durée de conservation"),
					("reorder_level", "Seuil de réapprovisionnement"),
					("reorder_qty", "Quantité à commander"),
				):
					if row[key] not in (None, ""):
						row[key] = as_number(row[key], sheet, number, report, label, positive=False)
				if row["reorder_qty"] and not row["reorder_level"]:
					report.error(sheet, number, "une quantité à commander demande un seuil")
			else:
				row["sizes"], row["colours"] = split_list(row["sizes"]), split_list(row["colours"])
				if not row["sizes"] or not row["colours"]:
					report.error(sheet, number, "au moins une taille et une couleur")
			if frappe.db.exists("Item", code):
				report.skip("Article", code)
			else:
				plan[kind].append(row)

	def item_exists(code):
		return code in codes_in_file or frappe.db.exists("Item", code)

	def stock_unit(code):
		return units_in_file.get(code) or frappe.db.get_value("Item", code, "stock_uom")

	boms = {}
	for number, row in data["Nomenclatures"]:
		product, component = str(row["product"]), str(row["component"])
		for code in (product, component):
			if not item_exists(code):
				report.error("Nomenclatures", number, f"article « {code} » introuvable")
		if row["uom"] and row["uom"] not in known_uoms:
			report.error("Nomenclatures", number, f"unité « {row['uom']} » inconnue")
		elif (
			row["uom"]
			and item_exists(component)
			and not unit_converts(component, row["uom"], stock_unit(component))
		):
			report.error(
				"Nomenclatures",
				number,
				f"« {component} » se compte en {stock_unit(component)} : rien ne convertit {row['uom']} "
				f"en {stock_unit(component)} (donnez la quantité en {stock_unit(component)})",
			)
		qty = as_number(row["qty"], "Nomenclatures", number, report, "Quantité")
		product_qty = as_number(row["product_qty"], "Nomenclatures", number, report, "Quantité produite")
		bom = boms.setdefault(
			product, {"product": product, "product_qty": product_qty, "rows": [], "line": number}
		)
		if bom["product_qty"] != product_qty:
			report.error(
				"Nomenclatures",
				number,
				f"« {product} » : la quantité produite doit être la même sur toutes ses lignes",
			)
		bom["rows"].append({"item_code": component, "qty": qty, "uom": row["uom"]})
	for product, bom in boms.items():
		if frappe.db.exists("BOM", {"item": product, "is_active": 1, "docstatus": 1}):
			report.skip("Nomenclature", product)
		else:
			plan["boms"].append(bom)

	warehouses = {
		frappe.db.get_value("Warehouse", name, "warehouse_name"): name
		for name in (
			settings.raw_materials_warehouse,
			settings.wip_warehouse,
			settings.finished_goods_warehouse,
		)
		if name
	}
	batch_items = {row["code"] for row in plan["items"] if row["batch"]}
	expiry_items = {row["code"] for row in plan["items"] if row["expiry"]}
	buying = {row["code"]: row["buying"] for row in plan["items"]}
	lots = {}  # lot number -> the item it belongs to in this file
	counted = set()  # (item, warehouse, lot) already in the file
	for number, row in data["Stock initial"]:
		code = str(row["code"])
		if not item_exists(code):
			report.error("Stock initial", number, f"article « {code} » introuvable")
			continue
		if row["warehouse"] not in warehouses:
			report.error(
				"Stock initial", number, f"entrepôt « {row['warehouse']} » inconnu ({', '.join(warehouses)})"
			)
			continue
		if frappe.db.exists(
			"Stock Ledger Entry",
			{"item_code": code, "warehouse": warehouses[row["warehouse"]], "is_cancelled": 0},
		):
			report.error(
				"Stock initial",
				number,
				f"« {code} » a déjà du stock dans « {row['warehouse']} » : "
				"corrigez-le dans ERPNext (Stock Reconciliation)",
			)
			continue
		has_batch = code in batch_items or cint(frappe.db.get_value("Item", code, "has_batch_no"))
		has_expiry = code in expiry_items or cint(frappe.db.get_value("Item", code, "has_expiry_date"))
		if has_batch and not row["batch"]:
			report.error("Stock initial", number, f"« {code} » est suivi par lot : indiquez le lot")
		lot = str(row["batch"]) if has_batch and row["batch"] else None

		if (code, row["warehouse"], lot) in counted:
			in_lot = f", lot {lot}" if lot else ""
			report.error("Stock initial", number, f"« {code} » compté deux fois ({row['warehouse']}{in_lot})")
			continue
		counted.add((code, row["warehouse"], lot))

		if lot:
			# ERPNext names a batch by its number alone: one number, one item
			owner = lots.setdefault(lot, code)
			if owner == code:
				owner = frappe.db.get_value("Batch", lot, "item") or code
			if owner != code:
				report.error(
					"Stock initial",
					number,
					f"lot « {lot} » : c'est déjà le lot de l'article « {owner} ». Un numéro de lot ne sert "
					f"qu'à un article : ajoutez-lui le code de l'article (ex. {lot}-{code})",
				)

		expiry = as_date(row["expiry"], "Stock initial", number, report)
		if has_expiry and row["batch"] and not expiry:
			report.error("Stock initial", number, f"« {code} » : indiquez la date d'expiration du lot")

		# ERPNext refuses an opening count without a value: the line's cost,
		# else the item's purchase price
		rate = None
		if row["rate"] not in (None, ""):
			rate = as_number(row["rate"], "Stock initial", number, report, "Coût unitaire")
		elif not flt(buying[code] if code in buying else frappe.db.get_value("Item", code, "valuation_rate")):
			report.error(
				"Stock initial",
				number,
				f"« {code} » : coût unitaire à remplir (l'article n'a pas de prix d'achat)",
			)
		plan["stock"].append(
			{
				"item_code": code,
				"warehouse": warehouses[row["warehouse"]],
				"qty": as_number(row["qty"], "Stock initial", number, report, "Quantité"),
				"valuation_rate": rate,
				"batch_no": lot,
				"expiry_date": expiry,
			}
		)
	return plan


def twice(seen, name):
	"""True when this name (case ignored) was already met in the sheet."""
	key = str(name).casefold()
	if key in seen:
		return True
	seen.add(key)
	return False


def create(plan, settings, report, opening_date=None):
	for row in plan["suppliers"]:
		doc = frappe.get_doc(
			{
				"doctype": "Supplier",
				"supplier_name": row["name"],
				"supplier_type": "Company",
				"mobile_no": normalize_phone(row["phone"]),
				"address_line1": row["address"],
				"city": row["city"],
				"country": "Algeria",
				"tax_id": row["nif"],
			}
		).insert(ignore_permissions=True)
		report.done("Fournisseur", doc.name)

	for row in plan["customers"]:
		phone = normalize_phone(row["phone"])
		wilaya = find_wilaya(row["wilaya"])
		doc = frappe.get_doc(
			{
				"doctype": "Customer",
				"customer_name": row["name"],
				"customer_type": "Company" if row["nif"] or row["rc"] else "Individual",
				"mobile_no": phone,
				"dz_phone": phone,
				"email_id": row["email"],
				"address_line1": row["address"],
				"city": row["city"],
				"state": frappe.db.get_value("DZ Wilaya", wilaya, "wilaya_name") if wilaya else None,
				"dz_wilaya": wilaya,
				"country": "Algeria",
				"dz_nif": row["nif"],
				"dz_rc": row["rc"],
				"dz_nis": row["nis"],
				"dz_ai": row["ai"],
			}
		).insert(ignore_permissions=True)
		report.done("Client", doc.name)

	for row in plan["items"]:
		doc = frappe.get_doc(
			{
				"doctype": "Item",
				"item_code": str(row["code"]),
				"item_name": row["name"],
				"item_group": row["group"],
				"stock_uom": row["uom"],
				"is_stock_item": 1,
				"include_item_in_manufacturing": 1,
				"valuation_rate": row["buying"] or 0,
				"standard_rate": row["selling"] or 0,
				"has_batch_no": cint(row["batch"]),
				"create_new_batch": cint(row["batch"]),
				"has_expiry_date": cint(row["expiry"]),
				"shelf_life_in_days": cint(row["shelf_life"]) if row["expiry"] else 0,
				"reorder_levels": reorder_rows(row, settings),
			}
		).insert(ignore_permissions=True)
		report.done("Article", doc.name)

	for row in plan["models"]:
		create_model(row, report)

	for bom in plan["boms"]:
		doc = frappe.get_doc(
			{
				"doctype": "BOM",
				"item": bom["product"],
				"company": settings.company,
				"quantity": bom["product_qty"],
				"is_active": 1,
				"is_default": 1,
				"rm_cost_as_per": "Valuation Rate",
				"items": [{k: v for k, v in r.items() if v} for r in bom["rows"]],
			}
		)
		doc.insert(ignore_permissions=True)
		doc.submit()
		report.done("Nomenclature", doc.name)

	if plan["stock"]:
		create_opening_stock(plan["stock"], settings, report, opening_date)


def reorder_rows(row, settings):
	"""The item's reorder level, in the warehouse its group stores it in
	(raw materials and packaging: Matières premières)."""
	if not row["reorder_level"]:
		return []
	warehouse = frappe.db.get_value(
		"Item Default", {"parent": row["group"], "company": settings.company}, "default_warehouse"
	) or (
		settings.finished_goods_warehouse
		if row["group"] in FINISHED_GROUPS
		else settings.raw_materials_warehouse
	)
	return [
		{
			"warehouse": warehouse,
			"warehouse_reorder_level": row["reorder_level"],
			"warehouse_reorder_qty": row["reorder_qty"] or 0,
			"material_request_type": "Purchase",
		}
	]


def create_model(row, report):
	"""A clothing model: one template item, one variant per size and colour."""
	from erpnext.controllers.item_variant import create_variant

	add_attribute_values("Taille", row["sizes"])
	add_attribute_values("Couleur", row["colours"])
	template = frappe.get_doc(
		{
			"doctype": "Item",
			"item_code": str(row["code"]),
			"item_name": row["name"],
			"item_group": row["group"],
			"stock_uom": row["uom"],
			"is_stock_item": 1,
			"has_variants": 1,
			"variant_based_on": "Item Attribute",
			# no price on the model itself: ERPNext refuses prices on template items,
			# each variant gets the selling price below
			"attributes": [{"attribute": "Taille"}, {"attribute": "Couleur"}],
		}
	).insert(ignore_permissions=True)
	report.done("Modèle", template.name)
	for size in row["sizes"]:
		for colour in row["colours"]:
			variant = create_variant(template.name, {"Taille": size, "Couleur": colour})
			# ERPNext names it "<model>-XL-BLE" (abbreviations): invoices show the words
			variant.item_name = variant.description = f"{row['name']} {size} {colour}"
			variant.standard_rate = row["selling"] or 0
			variant.insert(ignore_permissions=True)
			report.done("Variante", variant.name)


def add_attribute_values(attribute, values):
	"""Add sizes or colours the template list does not have yet."""
	doc = frappe.get_doc("Item Attribute", attribute)
	existing = {row.attribute_value.lower() for row in doc.item_attribute_values}
	used_abbrs = {row.abbr for row in doc.item_attribute_values}
	added = False
	for value in values:
		if value.lower() in existing:
			continue
		abbr = value.upper().replace(" ", "")[:3]
		suffix = 1
		while abbr in used_abbrs:
			abbr = f"{value.upper()[:2]}{suffix}"
			suffix += 1
		doc.append("item_attribute_values", {"attribute_value": value, "abbr": abbr})
		existing.add(value.lower())
		used_abbrs.add(abbr)
		added = True
	if added:
		doc.save(ignore_permissions=True)


def create_opening_stock(rows, settings, report, opening_date=None):
	"""Stock Reconciliations "Opening Stock" for the counted lines, MAX_LINES
	lines each, dated opening_date when given (else today)."""
	for row in rows:
		if row["batch_no"] and not frappe.db.exists("Batch", row["batch_no"]):
			frappe.get_doc(
				{
					"doctype": "Batch",
					"batch_id": row["batch_no"],
					"item": row["item_code"],
					"expiry_date": row["expiry_date"],
					# made before the count: ERPNext picks the oldest batch first
					# (FIFO by manufacturing date), so it must not look newer than
					# what is produced after the opening date
					"manufacturing_date": getdate(opening_date) if opening_date else None,
				}
			).insert(ignore_permissions=True)
			report.done("Lot", row["batch_no"])

	from dz_core.setup.accounts import opening_account

	account = opening_account(settings.company)
	for start in range(0, len(rows), MAX_LINES):
		doc = frappe.get_doc(
			{
				"doctype": "Stock Reconciliation",
				"company": settings.company,
				"purpose": "Opening Stock",
				"expense_account": account,
				"posting_date": getdate(opening_date) if opening_date else getdate(),
				"set_posting_time": 1 if opening_date else 0,
				# the count opens its day: what moves that day comes after it
				**({"posting_time": "00:00:00"} if opening_date else {}),
				"items": [
					{
						"item_code": row["item_code"],
						"warehouse": row["warehouse"],
						"qty": row["qty"],
						"valuation_rate": row["valuation_rate"]
						if row["valuation_rate"] is not None
						else frappe.db.get_value("Item", row["item_code"], "valuation_rate"),
						"use_serial_batch_fields": 1 if row["batch_no"] else 0,
						"batch_no": row["batch_no"],
					}
					for row in rows[start : start + MAX_LINES]
				],
			}
		)
		doc.insert(ignore_permissions=True)
		doc.submit()
		report.done("Stock initial", doc.name)
