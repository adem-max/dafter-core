"""The opening of the books from the accountant's closing trial balance
(SCF advice Q13; DECISIONS.md, "Compliance").

    bench --site <site> execute dz_core.imports.opening.make_opening_template \\
        --kwargs "{'path': '/tmp/ouverture.xlsx'}"
    bench --site <site> execute dz_core.imports.opening.import_opening \\
        --kwargs "{'path': '/tmp/ouverture.xlsx', 'opening_date': '2027-01-01'}"          # check only
    bench --site <site> execute dz_core.imports.opening.import_opening \\
        --kwargs "{'path': '/tmp/ouverture.xlsx', 'opening_date': '2027-01-01', 'dry_run': 0}"

Three sheets: "Balance" (account number, label, debit, credit: the closing
trial balance), "Clients ouverts" and "Fournisseurs ouverts" (the invoices
still due). Everything goes through the opening account 479, which must end
at exactly zero:

- customers' and suppliers' open invoices: opening invoices built by
  ERPNext's Opening Invoice Creation Tool (against 479), one per line, all or
  nothing; their total must equal the balance's receivable and payable lines;
- stock: already counted through the import workbook (an opening count on
  479), since a journal entry cannot post to a stock account; the balance's
  stock lines must equal what that count posted;
- every other line (fixed assets and their depreciation, equity, bank,
  cash, taxes, loans...): one Journal Entry "Écritures d'à-nouveaux"
  (is_opening), plus the line that brings 479 to zero. Existing fixed assets
  are entered by hand as assets "Is Existing Asset" (their cost and
  depreciation are in this entry).

Refused, with every reason listed, when an account is unknown, the balance
does not balance, or the invoices or the stock differ from it.
"""

import json

import frappe
from frappe import _
from frappe.utils import cint, flt, getdate

from dz_core.imports.workbook import Report, as_date, as_number, clean, resolve_path
from dz_core.setup.accounts import opening_account

BALANCE, CUSTOMERS, SUPPLIERS = "Balance", "Clients ouverts", "Fournisseurs ouverts"
SHEETS = {
	BALANCE: [
		("Compte", "number", True),
		("Libellé", "label", False),
		("Débit", "debit", False),
		("Crédit", "credit", False),
	],
	CUSTOMERS: [
		("Client", "party", True),
		("N° de facture", "invoice", True),
		("Date", "date", True),
		("Échéance", "due_date", False),
		("Reste dû TTC", "amount", True),
	],
	SUPPLIERS: [
		("Fournisseur", "party", True),
		("N° de facture", "invoice", True),
		("Date", "date", True),
		("Échéance", "due_date", False),
		("Reste dû TTC", "amount", True),
	],
}
EXAMPLES = {
	BALANCE: ["EXEMPLE 5121", "Banque BDL", 250000, 0],
	CUSTOMERS: ["EXEMPLE Épicerie El Baraka", "F-2026-118", "15/12/2026", "15/01/2027", 42000],
	SUPPLIERS: ["EXEMPLE Apiculteurs de Médéa", "FA-778", "20/12/2026", "20/01/2027", 18500],
}
README = [
	"Ouverture des comptes à partir de la balance de clôture du comptable",
	"",
	"Balance : un compte par ligne, avec son numéro SCF tel qu'il existe dans Daftar (ex. 5121, 10113, 40132).",
	"Clients ouverts / Fournisseurs ouverts : les factures encore dues, montant restant TTC.",
	"Le total des factures doit égaler les comptes clients (411...) et fournisseurs (401...) de la balance.",
	"Le stock est compté avant, avec le fichier d'import (feuille Stock) : ses comptes doivent égaler la balance.",
	"Les dates sont au format JJ/MM/AAAA. Les lignes qui commencent par EXEMPLE sont ignorées.",
]
TOLERANCE = 0.01


def make_opening_template(path):
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill

	wb = Workbook()
	readme = wb.active
	readme.title = "Lisez-moi"
	for line in README:
		readme.append([line])
	readme["A1"].font = Font(bold=True, size=14)
	readme.column_dimensions["A"].width = 120
	for name, columns in SHEETS.items():
		ws = wb.create_sheet(name)
		ws.append([header for header, _key, _required in columns])
		for cell, (_header, _key, required) in zip(ws[1], columns, strict=True):
			cell.font = Font(bold=True)
			if required:
				cell.fill = PatternFill("solid", fgColor="FFF2CC")
		ws.append(EXAMPLES[name])
		for cell in ws[2]:
			cell.font = Font(italic=True, color="808080")
		for index, (header, _key, _required) in enumerate(columns):
			ws.column_dimensions[chr(ord("A") + index)].width = max(16, len(header) + 2)
		ws.freeze_panes = "A2"
	wb.save(path)
	print(f"Workbook written: {path}. Yellow columns are required.")
	return path


def read_rows(wb, name, report):
	"""[(excel row number, {key: value})], examples and blank rows left out."""
	if name not in wb.sheetnames:
		return []
	columns = SHEETS[name]
	ws = wb[name]
	header = [str(cell.value or "").strip() for cell in ws[1]]
	missing = [text for text, _key, _required in columns if text not in header]
	if missing:
		report.error(name, 1, "colonnes introuvables : " + ", ".join(missing))
		return []
	positions = {key: header.index(text) for text, key, _required in columns}
	rows = []
	for number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
		row = {key: clean(values[index] if index < len(values) else None) for key, index in positions.items()}
		if not any(value not in (None, "") for value in row.values()):
			continue
		if str(next(iter(row.values())) or "").upper().startswith("EXEMPLE"):
			continue
		empty = [text for text, key, required in columns if required and row[key] in (None, "")]
		if empty:
			report.error(name, number, "à remplir : " + ", ".join(empty))
			continue
		rows.append((number, row))
	return rows


def import_opening(path, opening_date, dry_run=1):
	"""Check the workbook; with dry_run=0 and no errors, post the opening."""
	from openpyxl import load_workbook

	from dz_core.dz_core.doctype.dz_core_settings.dz_core_settings import get_settings

	dry_run = cint(dry_run)
	company = get_settings().company
	if not company:
		frappe.throw("Remplissez d'abord DZ Core Settings (dz_core.setup.new_client.prepare).")
	opening_date = getdate(opening_date)
	report = Report()
	wb = load_workbook(resolve_path(path), data_only=True)
	data = {name: read_rows(wb, name, report) for name in SHEETS}
	plan = check(company, opening_date, data, report)
	if report.errors or dry_run:
		result = report.as_dict()
		result["mode"] = "vérification seulement" if dry_run else "rien n'a été créé"
		result["would post"] = {
			"factures clients": len(plan["customers"]),
			"factures fournisseurs": len(plan["suppliers"]),
			"lignes d'à-nouveaux": len(plan["journal"]),
		}
		print(json.dumps(result, indent=1, ensure_ascii=False, default=str))
		return result
	post(company, opening_date, plan, report)
	frappe.db.commit()
	result = report.as_dict()
	print(json.dumps(result, indent=1, ensure_ascii=False, default=str))
	return result


def ledger(company, number):
	return frappe.db.get_value(
		"Account",
		{"company": company, "account_number": str(number).strip(), "is_group": 0, "disabled": 0},
		["name", "account_type"],
		as_dict=True,
	)


def posted(company, account, until):
	return flt(
		frappe.db.sql(
			"""SELECT SUM(debit - credit) FROM `tabGL Entry`
			WHERE company = %s AND account = %s AND is_cancelled = 0 AND posting_date <= %s""",
			(company, account, until),
		)[0][0]
	)


def check(company, opening_date, data, report):
	plan = {"customers": [], "suppliers": [], "journal": []}
	temporary = opening_account(company)
	sides = {"Receivable": 0.0, "Payable": 0.0}
	total = 0.0
	for number, row in data[BALANCE]:
		amount = as_number(row["debit"] or 0, BALANCE, number, report, "Débit", positive=False) - as_number(
			row["credit"] or 0, BALANCE, number, report, "Crédit", positive=False
		)
		account = ledger(company, row["number"])
		if not account:
			report.error(BALANCE, number, f"compte « {row['number']} » inconnu ou désactivé dans Daftar")
			continue
		if account.name == temporary:
			report.error(BALANCE, number, "le compte d'ouverture 479 ne figure pas dans la balance")
			continue
		total += amount
		if account.account_type in sides:
			sides[account.account_type] += amount
		elif account.account_type == "Stock":
			counted = posted(company, account.name, opening_date)
			if abs(counted - amount) > TOLERANCE:
				report.error(
					BALANCE,
					number,
					f"stock {row['number']} : la balance dit {amount:.2f}, le comptage d'ouverture a posté "
					f"{counted:.2f} (comptez le stock avec le fichier d'import d'abord)",
				)
		elif amount:
			plan["journal"].append(frappe._dict(account=account.name, amount=amount, number=row["number"]))
	if abs(total) > TOLERANCE:
		report.error(BALANCE, "-", f"la balance n'est pas équilibrée : débit - crédit = {total:.2f}")

	for sheet, party_type, side, sign in (
		(CUSTOMERS, "Customer", "Receivable", 1),
		(SUPPLIERS, "Supplier", "Payable", -1),
	):
		invoices = 0.0
		for number, row in data[sheet]:
			party = frappe.db.get_value(party_type, {"name": row["party"]}) or frappe.db.get_value(
				party_type, {frappe.scrub(party_type) + "_name": row["party"]}
			)
			if not party:
				report.error(sheet, number, f"« {row['party']} » n'existe pas (importez d'abord les tiers)")
				continue
			amount = as_number(row["amount"], sheet, number, report, "Reste dû")
			invoices += amount
			plan["customers" if party_type == "Customer" else "suppliers"].append(
				frappe._dict(
					party_type=party_type,
					party=party,
					invoice=str(row["invoice"]),
					date=as_date(row["date"], sheet, number, report),
					due_date=as_date(row["due_date"], sheet, number, report),
					amount=amount,
				)
			)
		if abs(invoices - sign * sides[side]) > TOLERANCE:
			report.error(
				sheet,
				"-",
				f"total des factures {invoices:.2f} différent des comptes {'clients' if sign > 0 else 'fournisseurs'}"
				f" de la balance ({sign * sides[side]:.2f})",
			)
	return plan


def post(company, opening_date, plan, report):
	temporary = opening_account(company)
	frappe.db.savepoint("dz_opening")
	for invoice_type, rows in (("Sales", plan["customers"]), ("Purchase", plan["suppliers"])):
		if not rows:
			continue
		tool = frappe.get_doc(
			{
				"doctype": "Opening Invoice Creation Tool",
				"company": company,
				"invoice_type": invoice_type,
				"invoices": [
					{
						"party_type": row.party_type,
						"party": row.party,
						# on the opening day (the year before may have no fiscal year
						# here); the invoice's own date and number are kept on it
						"posting_date": opening_date,
						"due_date": row.due_date or opening_date,
						"outstanding_amount": row.amount,
						"qty": 1,
						"item_name": _("Facture d'ouverture {0}").format(row.invoice),
						"temporary_opening_account": temporary,
						"supplier_invoice_date": row.date if invoice_type == "Purchase" else None,
					}
					for row in rows
				],
			}
		)
		for row, values in zip(rows, tool.get_invoices(), strict=True):
			values.pop("invoice_number", None)
			if invoice_type == "Purchase":
				values["bill_no"] = row.invoice
			else:
				values["remarks"] = _("Facture d'ouverture {0} du {1}").format(
					row.invoice, row.date.strftime("%d/%m/%Y") if row.date else "-"
				)
			doc = frappe.get_doc(values)
			doc.flags.ignore_mandatory = True
			doc.flags.dont_auto_add_taxes = True
			doc.insert()
			doc.submit()
			report.done(doc.doctype, doc.name)

	if plan["journal"]:
		entry = frappe.new_doc("Journal Entry")
		entry.update(
			{
				"company": company,
				"posting_date": opening_date,
				"voucher_type": "Opening Entry",
				"is_opening": "Yes",
				"user_remark": _("Écritures d'à-nouveaux au {0}").format(opening_date.strftime("%d/%m/%Y")),
			}
		)
		for line in plan["journal"]:
			entry.append(
				"accounts",
				{
					"account": line.account,
					"debit_in_account_currency": max(line.amount, 0),
					"credit_in_account_currency": max(-line.amount, 0),
				},
			)
		rest = -sum(line.amount for line in plan["journal"])
		entry.append(
			"accounts",
			{
				"account": temporary,
				"debit_in_account_currency": max(rest, 0),
				"credit_in_account_currency": max(-rest, 0),
			},
		)
		entry.insert()
		entry.submit()
		report.done("Journal Entry", entry.name)

	left = posted(company, temporary, opening_date)
	if abs(left) > TOLERANCE:
		frappe.db.rollback(save_point="dz_opening")
		frappe.throw(
			_("Le compte d'ouverture 479 ne tombe pas à zéro ({0}) : rien n'a été posté.").format(
				f"{left:.2f}"
			)
		)
